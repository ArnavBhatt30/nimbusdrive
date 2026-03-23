import os

def extract_text(file_bytes: bytes, filename: str) -> str:
    """
    Extract text content from uploaded files.
    Returns empty string if extraction fails or filetype unsupported.
    """
    ext = filename.split('.')[-1].lower()

    try:
        # Plain text files
        if ext in ['txt', 'md', 'csv', 'json', 'xml', 'html', 'py', 'js', 'ts', 'css']:
            return file_bytes.decode('utf-8', errors='ignore')[:5000]

        # PDF files
        if ext == 'pdf':
            return _extract_pdf(file_bytes)

        # Word documents
        if ext == 'docx':
            return _extract_docx(file_bytes)

        # Excel / CSV already handled above for csv
        if ext in ['xlsx', 'xls']:
            return _extract_excel(file_bytes)

    except Exception as e:
        print(f"Content extraction failed for {filename}: {e}")

    return ""


def _extract_pdf(file_bytes: bytes) -> str:
    try:
        import fitz  # PyMuPDF
        import io
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text = ""
        for page in doc:
            text += page.get_text()
            if len(text) > 5000:
                break
        return text[:5000]
    except ImportError:
        print("PyMuPDF not installed. Run: pip install pymupdf")
        return ""


def _extract_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
        import io
        doc = Document(io.BytesIO(file_bytes))
        text = "\n".join([para.text for para in doc.paragraphs])
        return text[:5000]
    except ImportError:
        print("python-docx not installed. Run: pip install python-docx")
        return ""


def _extract_excel(file_bytes: bytes) -> str:
    try:
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        text = ""
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join([str(cell) for cell in row if cell is not None])
                if row_text.strip():
                    text += row_text + "\n"
                if len(text) > 5000:
                    break
        return text[:5000]
    except ImportError:
        print("openpyxl not installed. Run: pip install openpyxl")
        return ""
